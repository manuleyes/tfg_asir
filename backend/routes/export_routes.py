"""
routes/export_routes.py - Exportar datos a CSV y Excel
"""
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from models.database import get_db
from models.person import Person
from models.vehicle import Vehicle
from models.alert import Alert
from models.camera import Camera
from middleware.auth_middleware import get_current_user
import csv
import io
import logging
from datetime import datetime
from typing import Optional

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/export", tags=["Export"])


def _make_csv_response(rows: list[dict], filename: str) -> StreamingResponse:
    """Crear StreamingResponse con contenido CSV"""
    if not rows:
        rows = [{}]
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=rows[0].keys())
    writer.writeheader()
    writer.writerows(rows)
    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _make_excel_response(rows: list[dict], filename: str, sheet_name: str = "Datos") -> StreamingResponse:
    """Crear StreamingResponse con contenido Excel"""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(
            status_code=501,
            detail="openpyxl no instalado. Ejecuta: pip install openpyxl",
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    if not rows:
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return StreamingResponse(
            iter([buffer.getvalue()]),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )

    # Cabeceras con estilo
    headers = list(rows[0].keys())
    header_fill = PatternFill("solid", fgColor="2563EB")
    header_font = Font(bold=True, color="FFFFFF")

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header.replace("_", " ").title())
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Datos
    for row_idx, row in enumerate(rows, 2):
        for col_idx, key in enumerate(headers, 1):
            value = row.get(key)
            # Convertir datetime a string para compatibilidad
            if isinstance(value, datetime):
                value = value.strftime("%Y-%m-%d %H:%M:%S")
            ws.cell(row=row_idx, column=col_idx, value=value)

    # Auto-ajustar ancho de columnas
    for col_idx, header in enumerate(headers, 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(
            len(str(header)),
            *(len(str(row.get(header, "") or "")) for row in rows),
        )
        ws.column_dimensions[col_letter].width = min(max_len + 4, 50)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        iter([buffer.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ─── PERSONAS ────────────────────────────────────────────────────────────────

@router.get("/personas/csv")
def export_personas_csv(
    limit: int = Query(1000, le=10000),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Exportar personas detectadas a CSV"""
    persons = db.query(Person).order_by(Person.detected_at.desc()).limit(limit).all()
    rows = [
        {
            "id": p.id,
            "camera_id": p.camera_id,
            "person_id": p.person_id,
            "confidence": p.confidence,
            "facial_features": p.facial_features,
            "times_detected": p.times_detected,
            "detected_at": p.detected_at,
            "image_path": p.image_path,
        }
        for p in persons
    ]
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    return _make_csv_response(rows, f"personas_{ts}.csv")


@router.get("/personas/excel")
def export_personas_excel(
    limit: int = Query(1000, le=10000),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Exportar personas detectadas a Excel"""
    persons = db.query(Person).order_by(Person.detected_at.desc()).limit(limit).all()
    rows = [
        {
            "id": p.id,
            "camera_id": p.camera_id,
            "person_id": p.person_id,
            "confidence": p.confidence,
            "facial_features": p.facial_features,
            "times_detected": p.times_detected,
            "detected_at": p.detected_at,
            "image_path": p.image_path,
        }
        for p in persons
    ]
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    return _make_excel_response(rows, f"personas_{ts}.xlsx", "Personas")


# ─── VEHÍCULOS ───────────────────────────────────────────────────────────────

@router.get("/vehiculos/csv")
def export_vehiculos_csv(
    limit: int = Query(1000, le=10000),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Exportar vehículos detectados a CSV"""
    vehicles = db.query(Vehicle).order_by(Vehicle.detected_at.desc()).limit(limit).all()
    rows = [
        {
            "id": v.id,
            "camera_id": v.camera_id,
            "license_plate": v.license_plate,
            "vehicle_model": v.vehicle_model,
            "vehicle_color": v.vehicle_color,
            "confidence": v.confidence,
            "times_detected": v.times_detected,
            "detected_at": v.detected_at,
        }
        for v in vehicles
    ]
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    return _make_csv_response(rows, f"vehiculos_{ts}.csv")


@router.get("/vehiculos/excel")
def export_vehiculos_excel(
    limit: int = Query(1000, le=10000),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Exportar vehículos detectados a Excel"""
    vehicles = db.query(Vehicle).order_by(Vehicle.detected_at.desc()).limit(limit).all()
    rows = [
        {
            "id": v.id,
            "camera_id": v.camera_id,
            "license_plate": v.license_plate,
            "vehicle_model": v.vehicle_model,
            "vehicle_color": v.vehicle_color,
            "confidence": v.confidence,
            "times_detected": v.times_detected,
            "detected_at": v.detected_at,
        }
        for v in vehicles
    ]
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    return _make_excel_response(rows, f"vehiculos_{ts}.xlsx", "Vehículos")


# ─── ALERTAS ─────────────────────────────────────────────────────────────────

@router.get("/alertas/csv")
def export_alertas_csv(
    severity: Optional[str] = Query(None),
    limit: int = Query(1000, le=10000),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Exportar alertas a CSV"""
    query = db.query(Alert)
    if severity:
        query = query.filter(Alert.severity == severity)
    alerts = query.order_by(Alert.created_at.desc()).limit(limit).all()
    rows = [
        {
            "id": a.id,
            "camera_id": a.camera_id,
            "alert_type": a.alert_type,
            "title": a.title,
            "description": a.description,
            "severity": a.severity,
            "is_read": a.is_read,
            "created_at": a.created_at,
        }
        for a in alerts
    ]
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    return _make_csv_response(rows, f"alertas_{ts}.csv")


@router.get("/alertas/excel")
def export_alertas_excel(
    severity: Optional[str] = Query(None),
    limit: int = Query(1000, le=10000),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Exportar alertas a Excel"""
    query = db.query(Alert)
    if severity:
        query = query.filter(Alert.severity == severity)
    alerts = query.order_by(Alert.created_at.desc()).limit(limit).all()
    rows = [
        {
            "id": a.id,
            "camera_id": a.camera_id,
            "alert_type": a.alert_type,
            "title": a.title,
            "description": a.description,
            "severity": a.severity,
            "is_read": a.is_read,
            "created_at": a.created_at,
        }
        for a in alerts
    ]
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    return _make_excel_response(rows, f"alertas_{ts}.xlsx", "Alertas")


# ─── CÁMARAS ─────────────────────────────────────────────────────────────────

@router.get("/camaras/csv")
def export_camaras_csv(
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Exportar cámaras a CSV"""
    cameras = db.query(Camera).all()
    rows = [
        {
            "id": c.id,
            "name": c.name,
            "url": c.url,
            "location": c.location,
            "description": c.description,
            "is_active": c.is_active,
            "created_at": c.created_at,
        }
        for c in cameras
    ]
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    return _make_csv_response(rows, f"camaras_{ts}.csv")


@router.get("/camaras/excel")
def export_camaras_excel(
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Exportar cámaras a Excel"""
    cameras = db.query(Camera).all()
    rows = [
        {
            "id": c.id,
            "name": c.name,
            "url": c.url,
            "location": c.location,
            "description": c.description,
            "is_active": c.is_active,
            "created_at": c.created_at,
        }
        for c in cameras
    ]
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    return _make_excel_response(rows, f"camaras_{ts}.xlsx", "Cámaras")
